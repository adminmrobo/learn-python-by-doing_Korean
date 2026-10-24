"""1편 실습 데이터(ep1-cafe.zip)를 만든다. 값은 모두 연습용으로 지어낸 것이다.
usage: python tools/make_ep1_data.py"""
import io, json, os, zipfile
import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
rng = np.random.default_rng(2025)

MENU = {  # menu: (price, cost, category)
    "Americano": (4500, 1200, "Drink"), "Latte": (5000, 1600, "Drink"), "Iced Tea": (5500, 1100, "Drink"),
    "Cheesecake": (6000, 2400, "Dessert"), "Croissant": (3500, 1300, "Dessert"),
}

# ---------- koi_sales.csv : 2025년 영수증 줄 ----------
rows = []
for d in pd.date_range("2025-01-01", "2025-12-31"):
    lam = 34 + (14 if d.dayofweek >= 5 else 0) + (8 if d.month in (7, 8) else 0)   # 하루 영수증 줄 수
    summer = d.month in (6, 7, 8, 9)
    p = np.array([.32, .24, .16 if summer else .06, .12, .16 if summer else .26]); p = p / p.sum()
    for k in range(rng.poisson(lam)):
        m = rng.choice(list(MENU), p=p)
        qty = int(rng.choice([1, 1, 1, 2, 2, 3]))
        cust = rng.choice(["Regular", "Walk-in", "Tourist", None], p=[.38, .45, .12, .05])
        disc = 500 * qty if cust == "Regular" and rng.random() < .5 else 0
        rows.append([d, m, qty, MENU[m][0], disc, cust])
sales = pd.DataFrame(rows, columns=["date", "menu", "qty", "price", "discount", "customer"])
sales.insert(0, "receipt_id", [int(d.strftime("%y%m%d")) * 1000 + i for i, d in
                               zip(sales.groupby("date").cumcount() + 1, sales["date"])])
# 함정: 같은 영수증이 두 번 찍힌 줄 3개 (1-12 검증에서 잡힌다)
dup = sales.iloc[[4120, 9051, 13007]]
sales = pd.concat([sales, dup]).sort_values(["date", "receipt_id"], kind="stable").reset_index(drop=True)
sales["date"] = sales["date"].dt.strftime("%Y-%m-%d")

# ---------- koi_cost.xlsx : 지저분한 원가표 ----------
wb = Workbook(); ws = wb.active; ws.title = "cost"
ws.append(["카페 Koi 메뉴 원가표"]); ws.merge_cells("A1:D1")
ws.append(["2025년 기준 · 단위: 원"]); ws.merge_cells("A2:D2")
ws.append([" menu", "price ", " cost ", "updated"])          # 열 이름 앞뒤에 공백
data = [("Americano", "2025-01-02"), ("Latte ", "2025-01-02"), (None, None), ("Iced Tea", "2025-05-20"),
        ("Cheesecake", "2025-03-11"), ("Croissant ", "-")]   # Latte, Croissant 끝에 공백 / 빈 행 / '-'
for name, upd in data:
    if name is None:
        ws.append([None, None, None, None]); continue
    pr, co, _ = MENU[name.strip()]
    ws.append([name, f"₩{pr:,}", f"₩{co:,}", upd])
buf = io.BytesIO(); wb.save(buf); cost_xlsx = buf.getvalue()

# ---------- koi_delivery.json : 배달앱 주문 응답 (7월) ----------
orders = []
for i in range(24):
    d = pd.Timestamp("2025-07-01") + pd.Timedelta(days=int(rng.integers(0, 31)))
    m = rng.choice(["Americano", "Latte", "Iced Tea", "Cheesecake"], p=[.3, .25, .3, .15])
    orders.append({"order_id": f"D-{2501 + i}", "date": d.strftime("%Y-%m-%d"), "menu": m,
                   "qty": int(rng.choice([1, 2, 2, 3])), "price": MENU[m][0]})
orders.sort(key=lambda o: o["date"])
delivery = {"status": "ok", "store": "Cafe Koi", "month": "2025-07", "count": len(orders), "orders": orders}

# ---------- koi_fixed.csv : 월별 고정비 ----------
fx = []
for m in range(1, 13):
    fx += [[m, "Rent", 1_800_000], [m, "Staff", 3_200_000 + (400_000 if m in (7, 8, 12) else 0)],
           [m, "Utilities", int(round(280_000 + (170_000 if m in (1, 2, 7, 8) else 0) + rng.integers(0, 40) * 1000, -3))]]
fixed = pd.DataFrame(fx, columns=["month", "item", "amount"])

# ---------- koi_stock.csv : 4월 원두 재고 변동 (kg) ----------
st = [["2025-04-01", "opening", 8.0]]
for day, kg in [(3, 10.0), (14, 10.0), (24, 6.0)]:
    st.append([f"2025-04-{day:02d}", "in", kg])
for day in range(1, 31):
    st.append([f"2025-04-{day:02d}", "used", round(float(rng.uniform(0.75, 0.95)), 1)])
st += [["2025-04-18", "waste", 0.8], ["2025-04-29", "waste", 0.6]]
stock = pd.DataFrame(st, columns=["date", "type", "kg"]).sort_values("date", kind="stable")
closing = round(8.0 + 26.0 - stock.loc[stock.type == "used", "kg"].sum() - 1.4, 1)
stock = pd.concat([stock, pd.DataFrame([["2025-04-30", "count", closing]], columns=stock.columns)])

out = os.path.join(ROOT, "src", "data", "ep1-cafe.zip")
with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
    for name, df in [("koi_sales.csv", sales), ("koi_fixed.csv", fixed), ("koi_stock.csv", stock)]:
        z.writestr(name, "﻿" + df.to_csv(index=False))
    z.writestr("koi_cost.xlsx", cost_xlsx)
    z.writestr("koi_delivery.json", json.dumps(delivery, ensure_ascii=False, indent=2))
print("wrote", out, "| sales rows", len(sales), "| stock closing", closing)
